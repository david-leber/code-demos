"""Create sample resume files for testing."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt


def create_sample_word_resume_1():
    """Create a simple Word resume sample."""
    doc = Document()

    # Name
    name = doc.add_paragraph("John Smith")
    name.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    name.runs[0].bold = True
    name.runs[0].font.size = Pt(16)

    # Contact
    contact = doc.add_paragraph("john.smith@email.com | (555) 123-4567 | linkedin.com/in/johnsmith")
    contact.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_paragraph()  # Spacing

    # Education
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("Bachelor of Science in Computer Science")
    doc.add_paragraph("University of California, Berkeley | 2018-2022")
    doc.add_paragraph("GPA: 3.8/4.0")

    doc.add_paragraph()

    # Experience
    doc.add_heading("EXPERIENCE", level=1)

    doc.add_paragraph("Software Engineer | Tech Company Inc. | 2022-Present")
    bullets = [
        "• Developed and maintained web applications using Python and React",
        "• Collaborated with cross-functional teams to deliver features",
        "• Improved system performance by 40% through optimization"
    ]
    for bullet in bullets:
        doc.add_paragraph(bullet)

    doc.add_paragraph()

    # Skills
    doc.add_heading("SKILLS", level=1)
    doc.add_paragraph("Programming: Python, JavaScript, Java, C++")
    doc.add_paragraph("Frameworks: React, Django, Flask, Node.js")
    doc.add_paragraph("Tools: Git, Docker, AWS, Linux")

    return doc


def create_sample_word_resume_2():
    """Create a marketing resume sample."""
    doc = Document()

    # Name
    name = doc.add_paragraph("Sarah Johnson")
    name.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    name.runs[0].bold = True
    name.runs[0].font.size = Pt(18)

    # Contact
    contact = doc.add_paragraph("sarah.j@example.com | 555-987-6543")
    contact.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_paragraph()

    # Professional Summary
    doc.add_heading("PROFESSIONAL SUMMARY", level=1)
    doc.add_paragraph(
        "Results-driven Marketing Manager with 5+ years of experience in digital marketing, "
        "brand strategy, and campaign management. Proven track record of increasing brand "
        "awareness and driving revenue growth."
    )

    doc.add_paragraph()

    # Work Experience
    doc.add_heading("WORK EXPERIENCE", level=1)

    doc.add_paragraph("Marketing Manager | Digital Solutions Corp | 2020-Present")
    bullets = [
        "- Led marketing campaigns that increased customer acquisition by 60%",
        "- Managed a team of 5 marketing specialists",
        "- Developed and executed social media strategy reaching 500K+ followers",
        "- Analyzed market trends and competitor activities to inform strategy"
    ]
    for bullet in bullets:
        doc.add_paragraph(bullet)

    doc.add_paragraph()
    doc.add_paragraph("Marketing Coordinator | StartUp Inc | 2018-2020")
    bullets2 = [
        "* Created content for email campaigns with 25% open rate",
        "* Managed company blog and increased organic traffic by 80%",
        "* Coordinated trade show events and conferences"
    ]
    for bullet in bullets2:
        doc.add_paragraph(bullet)

    doc.add_paragraph()

    # Education
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("MBA in Marketing | Stanford University | 2018")
    doc.add_paragraph("BA in Communications | UCLA | 2016")

    return doc


def create_sample_word_resume_3():
    """Create a data science resume sample."""
    doc = Document()

    # Name
    name = doc.add_paragraph("Dr. Emily Chen")
    name.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    name.runs[0].bold = True
    name.runs[0].font.size = Pt(14)

    # Contact
    contact = doc.add_paragraph("emily.chen@research.edu | github.com/echen")
    contact.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_paragraph()

    # Skills
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph("Languages: Python, R, SQL, Scala")
    doc.add_paragraph("ML/AI: TensorFlow, PyTorch, scikit-learn, Keras")
    doc.add_paragraph("Big Data: Spark, Hadoop, Hive")
    doc.add_paragraph("Visualization: Tableau, Matplotlib, D3.js")

    doc.add_paragraph()

    # Experience
    doc.add_heading("Professional Experience", level=1)

    doc.add_paragraph("Senior Data Scientist | AI Research Lab | 2021-Present")
    bullets = [
        "1. Developed machine learning models achieving 95% accuracy",
        "2. Led research initiatives in natural language processing",
        "3. Published 5 papers in top-tier conferences (NeurIPS, ICML)",
        "4. Mentored junior data scientists and interns"
    ]
    for bullet in bullets:
        doc.add_paragraph(bullet)

    doc.add_paragraph()

    # Projects
    doc.add_heading("Notable Projects", level=1)
    doc.add_paragraph("Predictive Analytics Platform: Built end-to-end ML pipeline processing 10M+ records/day")
    doc.add_paragraph("Customer Segmentation: Developed clustering algorithm reducing marketing costs by $2M")

    doc.add_paragraph()

    # Education
    doc.add_heading("Education", level=1)
    doc.add_paragraph("Ph.D. in Computer Science | MIT | 2021")
    doc.add_paragraph("Specialization: Machine Learning & Artificial Intelligence")
    doc.add_paragraph("M.S. in Statistics | Carnegie Mellon University | 2018")

    return doc


def save_sample_resumes():
    """Create and save all sample resumes."""
    examples_dir = Path("examples")
    examples_dir.mkdir(exist_ok=True)

    # Create Word resumes
    doc1 = create_sample_word_resume_1()
    doc1.save(examples_dir / "sample_resume_tech.docx")
    print("Created: sample_resume_tech.docx")

    doc2 = create_sample_word_resume_2()
    doc2.save(examples_dir / "sample_resume_marketing.docx")
    print("Created: sample_resume_marketing.docx")

    doc3 = create_sample_word_resume_3()
    doc3.save(examples_dir / "sample_resume_data_science.docx")
    print("Created: sample_resume_data_science.docx")


if __name__ == "__main__":
    save_sample_resumes()
